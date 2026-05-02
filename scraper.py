"""
Financial data scraper — CI Financial, Tangerine, Highgate Group.

Usage:
    python3 scraper.py                  # scrape all sites
    python3 scraper.py ci               # scrape only CI Financial
    python3 scraper.py tangerine        # scrape only Tangerine
    python3 scraper.py highgate         # scrape only Highgate
    python3 scraper.py --headless       # run without visible browser

Output: financial_data_YYYY-MM-DD.xlsx
"""

import sys
import os
import re
from datetime import date
from dotenv import load_dotenv
from playwright.sync_api import sync_playwright
import pandas as pd

from scrapers.base import new_browser_context, AccountData
from scrapers import ci_financial, tangerine, highgate

load_dotenv()

SCRAPERS = {
    "ci":        ci_financial.scrape,
    "tangerine": tangerine.scrape,
    "highgate":  highgate.scrape,
}

SITE_URLS = {
    "ci":        "https://www.cifinancial.com/ci-gam/ca/en/index.html",
    "tangerine": "https://www.tangerine.ca/app/#/login/login-id?locale=en_CA",
    "highgate":  "https://highgategroup.ca/",
}


def run(targets: list[str], headless: bool = False) -> list[AccountData]:
    results = []
    with sync_playwright() as p:
        for key in targets:
            print(f"\n{'='*50}")
            print(f" Scraping: {key.upper()}")
            print(f"{'='*50}")
            browser, context = new_browser_context(p, headless=headless)
            page = context.new_page()
            try:
                data = SCRAPERS[key](page)
                results.append(data)
                if data.error:
                    print(f"[{key}] ERROR: {data.error}")
                else:
                    print(f"[{key}] Balances: {len(data.balances)} | Transactions: {len(data.transactions)}")
            finally:
                context.close()
                browser.close()
    return results


def _flatten_row(row_val) -> str:
    if isinstance(row_val, list):
        return " | ".join(str(v) for v in row_val if v)
    return str(row_val)


def write_excel(results: list[AccountData], output_path: str):
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        # Summary sheet
        summary_rows = []
        for d in results:
            summary_rows.append({
                "Source":        d.source,
                "Balances Found":     len(d.balances),
                "Transactions Found": len(d.transactions),
                "Error":         d.error or "",
            })
        pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Summary", index=False)
        _style_sheet(writer, "Summary")

        for d in results:
            sheet_prefix = re.sub(r"[^A-Za-z0-9]", "", d.source)[:12]

            # Balances sheet
            if d.balances:
                bal_df = pd.DataFrame(d.balances)
                bal_df.to_excel(writer, sheet_name=f"{sheet_prefix}_Balances", index=False)
                _style_sheet(writer, f"{sheet_prefix}_Balances")

            # Transactions sheet
            if d.transactions:
                tx_rows = []
                for tx in d.transactions:
                    flat = {k: (_flatten_row(v) if k == "Row" else v) for k, v in tx.items()}
                    tx_rows.append(flat)
                tx_df = pd.DataFrame(tx_rows)
                tx_df.to_excel(writer, sheet_name=f"{sheet_prefix}_Transactions", index=False)
                _style_sheet(writer, f"{sheet_prefix}_Transactions")

    print(f"\nSaved: {output_path}")


def _style_sheet(writer: pd.ExcelWriter, sheet_name: str):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    ws = writer.sheets[sheet_name]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for col_idx, col in enumerate(ws.columns, 1):
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

    ws.freeze_panes = "A2"


if __name__ == "__main__":
    args = sys.argv[1:]
    headless = "--headless" in args
    args = [a for a in args if a != "--headless"]

    targets = [a.lower() for a in args if a.lower() in SCRAPERS]
    if not targets:
        targets = list(SCRAPERS.keys())

    print(f"Targets: {', '.join(targets)} | Headless: {headless}")
    results = run(targets, headless=headless)

    output_file = f"financial_data_{date.today()}.xlsx"
    write_excel(results, output_file)
